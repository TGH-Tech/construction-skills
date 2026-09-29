#!/usr/bin/env python3
"""Deterministic critical-path (CPM) calculator for construction schedules.

It never decides scope, activities or durations — those are inputs — it only
computes their consequences: early/late start and finish, total float, the
connected critical chain and near-critical activities. Standard library only;
XLSX output needs openpyxl (`pip install openpyxl`), everything else does not.

Input (JSON):

    {
      "project": "Villa at Plot 12",
      "activities": [
        {"id": "A010", "name": "Mobilization", "duration": 5, "predecessors": []},
        {"id": "A020", "name": "Site establishment", "duration": 7,
         "predecessors": [{"id": "A010", "type": "FS", "lag": 0}]}
      ]
    }

  * duration  working days (> 0; 0 only with "milestone": true)
  * type      FS | SS | FF | SF (default FS)
  * lag       working days, negative for overlap (default 0)
  * a predecessor may also be written as a string: "A010", "A010 SS", "A010 FS+2"
  * source    optional — where the duration came from ("library (provisional)",
              "proposed, accepted", "user-supplied"); shown in CSV and XLSX
  * milestone optional — true allows a zero duration

A CSV with columns id,name,duration,predecessors (predecessors like
"A010; B020 SS+1") is accepted too.

Usage:

    python3 cpm.py schedule.json                  # table
    python3 cpm.py schedule.json --json           # full result as JSON
    python3 cpm.py schedule.json --csv out.csv
    python3 cpm.py schedule.json --xlsx out.xlsx --start 2026-10-05 \
        --workdays 1-6 --holiday 2026-10-20 --deadline 2027-10-04

Day 0 is the morning of the start date; an offset of N means N working days
have elapsed. Exit code 2 means the network cannot be calculated (a cycle, an
unknown predecessor, a missing duration) — the messages say what to fix.

Sample of Bleno's critical-path workflow. The full workflow — complete planning
library, mapping from your agreement, review cards, PDF and Excel — runs at
https://constructions.bleno.io; pilots at https://bleno.io.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
from collections import deque

RELATIONS = ("FS", "SS", "FF", "SF")


# ── input ────────────────────────────────────────────────────────────────

def _dep_from_text(text):
    m = re.fullmatch(r"\s*([\w.-]+)\s*(FS|SS|FF|SF)?\s*([+-]\s*\d+(?:\.\d+)?)?\s*", text, re.I)
    if not m:
        raise ValueError(f'cannot read predecessor "{text}" — write it like "A010", "A010 SS" or "A010 FS+2"')
    lag = float(m.group(3).replace(" ", "")) if m.group(3) else 0
    lag = int(lag) if float(lag).is_integer() else lag
    return {"id": m.group(1), "type": (m.group(2) or "FS").upper(), "lag": lag}


def _normalise(raw):
    acts = []
    for a in raw:
        preds = []
        for p in a.get("predecessors") or []:
            if isinstance(p, str):
                preds.append(_dep_from_text(p))
            else:
                preds.append({
                    "id": str(p.get("id") or p.get("predecessorId") or p.get("on")),
                    "type": str(p.get("type") or p.get("relation") or "FS").upper(),
                    "lag": p.get("lag", 0),
                })
        acts.append({
            "id": str(a["id"]),
            "name": str(a.get("name") or a["id"]),
            "duration": a.get("duration"),
            "milestone": bool(a.get("milestone")),
            "predecessors": preds,
            # where the duration came from, carried to every output
            "source": str(a.get("source") or ""),
        })
    return acts


def load(path):
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        raw = [{
            "id": r["id"],
            "name": r.get("name") or r["id"],
            "duration": float(r["duration"]) if r.get("duration", "").strip() else None,
            "predecessors": [p for p in re.split(r"[;,]", r.get("predecessors") or "") if p.strip()],
        } for r in rows]
        return "", _normalise(raw)
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return "", _normalise(data)
    return str(data.get("project") or ""), _normalise(data["activities"])


# ── validation ───────────────────────────────────────────────────────────

def _topological(acts):
    ids = {a["id"] for a in acts}
    indeg = {a["id"]: 0 for a in acts}
    succ = {}
    for a in acts:
        for p in a["predecessors"]:
            if p["id"] not in ids:
                continue
            indeg[a["id"]] += 1
            succ.setdefault(p["id"], []).append(a["id"])
    queue = deque(a["id"] for a in acts if indeg[a["id"]] == 0)
    order = []
    while queue:
        i = queue.popleft()
        order.append(i)
        for n in succ.get(i, []):
            indeg[n] -= 1
            if indeg[n] == 0:
                queue.append(n)
    return order if len(order) == len(acts) else None


def _find_cycle(acts):
    by_id = {a["id"]: a for a in acts}
    state, stack, found = {}, [], []

    def visit(i):
        nonlocal found
        if state.get(i) == "done":
            return False
        if state.get(i) == "visiting":
            found = stack[stack.index(i):] + [i]
            return True
        state[i] = "visiting"
        stack.append(i)
        for p in by_id[i]["predecessors"]:
            if p["id"] in by_id and visit(p["id"]):
                return True
        stack.pop()
        state[i] = "done"
        return False

    for a in acts:
        if visit(a["id"]):
            break
    return list(reversed(found))


def validate(acts):
    """Returns (blocking, warnings) — lists of messages."""
    blocking, warnings = [], []
    seen = set()
    by_id = {}
    for a in acts:
        if a["id"] in seen:
            blocking.append(f'Activity id "{a["id"]}" appears more than once.')
        seen.add(a["id"])
        by_id[a["id"]] = a
    for a in acts:
        d = a["duration"]
        if not a["milestone"] and not (isinstance(d, (int, float)) and d > 0):
            blocking.append(f'"{a["name"]}" ({a["id"]}) has no usable duration. How many working days should it take?')
        for p in a["predecessors"]:
            if p["id"] not in by_id:
                blocking.append(f'"{a["name"]}" ({a["id"]}) depends on "{p["id"]}", which is not in the schedule.')
            if p["type"] not in RELATIONS:
                blocking.append(f'"{a["name"]}" ({a["id"]}) uses relationship "{p["type"]}"; only FS, SS, FF and SF are supported.')
            if not isinstance(p["lag"], (int, float)):
                blocking.append(f'"{a["name"]}" ({a["id"]}) has a non-numeric lag on its link from {p["id"]}.')
    if blocking:
        return blocking, warnings
    if _topological(acts) is None:
        cycle = _find_cycle(acts)
        return [f'The schedule contains a circular dependency: {" → ".join(cycle)}. CPM cannot be calculated until this is corrected.'], warnings
    if acts:
        if not any(not a["predecessors"] for a in acts):
            blocking.append("No start activity — every activity has a predecessor.")
        has_succ = {p["id"] for a in acts for p in a["predecessors"]}
        terminal = [a for a in acts if a["id"] not in has_succ]
        if not terminal:
            blocking.append("No finish activity — every activity feeds another one.")
        elif len(terminal) > 1:
            names = ", ".join(f'"{a["name"]}"' for a in terminal)
            warnings.append(f"{len(terminal)} activities have nothing following them: {names}. Only the project's finish should end the chain — link the rest or confirm they really are ends.")
        # connectivity, links treated as undirected
        nbr = {a["id"]: set() for a in acts}
        for a in acts:
            for p in a["predecessors"]:
                nbr[a["id"]].add(p["id"])
                nbr[p["id"]].add(a["id"])
        seen, queue = {acts[0]["id"]}, deque([acts[0]["id"]])
        while queue:
            for n in nbr[queue.popleft()]:
                if n not in seen:
                    seen.add(n)
                    queue.append(n)
        loose = [a["id"] for a in acts if a["id"] not in seen]
        if loose:
            warnings.append(f"{len(loose)} activit{'y is' if len(loose) == 1 else 'ies are'} not connected to the main network: {', '.join(loose)}.")
    return blocking, warnings


# ── calculation ──────────────────────────────────────────────────────────

def _es_from(p, pred, duration):
    t, lag = p["type"], p["lag"]
    return {"FS": pred["ef"] + lag, "SS": pred["es"] + lag,
            "FF": pred["ef"] + lag - duration, "SF": pred["es"] + lag - duration}[t]


def _lf_from(p, succ, pred_duration):
    t, lag = p["type"], p["lag"]
    return {"FS": succ["ls"] - lag, "SS": succ["ls"] - lag + pred_duration,
            "FF": succ["lf"] - lag, "SF": succ["lf"] - lag + pred_duration}[t]


def _deps_text(preds):
    out = []
    for p in preds:
        lag = p["lag"]
        out.append(p["id"] + ("" if p["type"] == "FS" and not lag else f" {p['type']}") +
                   (f"{lag:+g}" if lag else ""))
    return "; ".join(out)


def calculate(acts):
    blocking, warnings = validate(acts)
    if blocking:
        raise ValueError(" ".join(blocking))
    order = _topological(acts)
    by_id = {a["id"]: a for a in acts}
    dur = {a["id"]: (a["duration"] or 0) for a in acts}
    rows = {}
    for i in order:
        a = by_id[i]
        es = 0 if not a["predecessors"] else max(0, *[_es_from(p, rows[p["id"]], dur[i]) for p in a["predecessors"]])
        rows[i] = {"id": i, "name": a["name"], "duration": dur[i], "es": es, "ef": es + dur[i],
                   "source": a["source"], "predecessors": _deps_text(a["predecessors"])}
    successors = {}
    for a in acts:
        for p in a["predecessors"]:
            successors.setdefault(p["id"], []).append((p, a["id"]))
    finish = max([r["ef"] for r in rows.values()] or [0])
    for i in reversed(order):
        r = rows[i]
        succ = successors.get(i, [])
        r["lf"] = min(finish, *[_lf_from(p, rows[s], dur[i]) for p, s in succ]) if succ else finish
        r["ls"] = r["lf"] - dur[i]
        r["float"] = r["ls"] - r["es"]
        r["critical"] = abs(r["float"]) < 1e-9

    # the critical path is a CONNECTED chain: walk back from the finishing
    # activity through whichever critical predecessor actually drives its start
    chain = []
    ends = sorted((r["id"] for r in rows.values() if r["critical"] and r["ef"] == finish),
                  key=lambda i: bool(successors.get(i)))  # prefer a true end of the network
    current, guard = (ends[0] if ends else None), set()
    while current and current not in guard:
        guard.add(current)
        chain.append(current)
        row = rows[current]
        driving = [p for p in by_id[current]["predecessors"]
                   if rows[p["id"]]["critical"] and _es_from(p, rows[p["id"]], dur[current]) == row["es"]]
        current = driving[0]["id"] if driving else None
    chain.reverse()

    return {
        "finish_days": finish,
        "rows": [rows[i] for i in order],
        "critical_path": chain,
        # zero-float activities running alongside the chain (parallel critical branches)
        "also_critical": [i for i in order if rows[i]["critical"] and i not in chain],
        "near_critical": [i for i in order if 0 < rows[i]["float"] <= 2],
        "warnings": warnings,
    }


# ── calendar ─────────────────────────────────────────────────────────────

def _working(day, workdays, holidays):
    return day.isoweekday() in workdays and day.isoformat() not in holidays


def to_date(offset, start, workdays, holidays, kind):
    """Day 0 = morning of `start`. Starts land ON the next working day;
    finishes land on the working day that completes the offset."""
    d = start
    if kind == "start":
        remaining = offset
        while remaining > 0 or not _working(d, workdays, holidays):
            if _working(d, workdays, holidays):
                remaining -= 1
            d += dt.timedelta(days=1)
        return d
    remaining = max(1, offset)
    while remaining > 0:
        if _working(d, workdays, holidays):
            remaining -= 1
        if remaining > 0:
            d += dt.timedelta(days=1)
    return d


# ── output ───────────────────────────────────────────────────────────────

def _fmt(n):
    return str(int(n)) if float(n).is_integer() else f"{n:g}"


def _by_start(result):
    """Rows as a programme reads them: by early start, then id."""
    return sorted(result["rows"], key=lambda r: (r["es"], r["id"]))


def table(project, result, dates):
    head = ["ID", "Activity", "Dur", "ES", "EF", "LS", "LF", "Float", ""]
    if dates:
        head[7:7] = ["Start", "Finish"]
    lines = []
    for r in _by_start(result):
        cells = [r["id"], r["name"][:48], _fmt(r["duration"]), _fmt(r["es"]), _fmt(r["ef"]),
                 _fmt(r["ls"]), _fmt(r["lf"]), _fmt(r["float"]), "CRITICAL" if r["critical"] else ""]
        if dates:
            s, f = dates[r["id"]]
            cells[7:7] = [s, f]
        lines.append(cells)
    widths = [max(len(str(x)) for x in col) for col in zip(head, *lines)]
    out = []
    if project:
        out.append(project)
    out.append("  ".join(h.ljust(w) for h, w in zip(head, widths)))
    out.append("  ".join("-" * w for w in widths))
    out += ["  ".join(str(c).ljust(w) for c, w in zip(row, widths)) for row in lines]
    out.append("")
    out.append(f"Project duration: {_fmt(result['finish_days'])} working days")
    out.append("Critical path: " + " → ".join(result["critical_path"]))
    if result["also_critical"]:
        out.append("Also critical, in parallel: " + ", ".join(result["also_critical"]))
    if result["near_critical"]:
        out.append("Near-critical (float 1–2 days): " + ", ".join(result["near_critical"]))
    for w in result["warnings"]:
        out.append("Check: " + w)
    for n in result.get("notes", []):
        out.append("Note: " + n)
    return "\n".join(out)


def write_csv(path, result, dates):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "name", "duration", "predecessors", "duration_source", "es", "ef", "ls", "lf",
                    "total_float", "critical", "start", "finish"])
        for r in _by_start(result):
            s, fi = dates.get(r["id"], ("", "")) if dates else ("", "")
            w.writerow([r["id"], r["name"], r["duration"], r["predecessors"], r["source"], r["es"], r["ef"],
                        r["ls"], r["lf"], r["float"], "yes" if r["critical"] else "", s, fi])


def write_xlsx(path, project, result, dates, day_labels=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("XLSX needs openpyxl: pip install openpyxl (or use --csv)")
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    head = ["ID", "Activity", "Days", "Predecessors", "Duration source", "ES", "EF", "LS", "LF",
            "Total float", "Critical", "Start", "Finish"]
    ws.append([project or "Critical-path schedule"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"{_fmt(result['finish_days'])} working days · critical path: " + " → ".join(result["critical_path"])])
    ws.append([])
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="161615")
    red = PatternFill("solid", fgColor="FBE3E1")
    ordered = _by_start(result)
    for r in ordered:
        s, f = dates.get(r["id"], ("", "")) if dates else ("", "")
        ws.append([r["id"], r["name"], r["duration"], r["predecessors"], r["source"], r["es"], r["ef"],
                   r["ls"], r["lf"], r["float"], "yes" if r["critical"] else "", s, f])
        if r["critical"]:
            for c in ws[ws.max_row]:
                c.fill = red
    # a simple bar per working day, so the file reads as a Gantt at a glance
    first_day_col = len(head) + 2
    days = int(result["finish_days"])
    labels = day_labels or [str(d + 1) for d in range(days)]
    for d in range(days):
        cell = ws.cell(row=4, column=first_day_col + d, value=labels[d])
        cell.font = Font(size=8)
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(first_day_col + d)].width = 5 if day_labels else 3
    bar, crit = PatternFill("solid", fgColor="8FA3B8"), PatternFill("solid", fgColor="C0392B")
    for i, r in enumerate(ordered, start=5):
        for d in range(int(r["es"]), int(r["ef"])):
            ws.cell(row=i, column=first_day_col + d).fill = crit if r["critical"] else bar
    for col, width in zip("ABCDEFGHIJKLM", [8, 38, 6, 16, 18, 6, 6, 6, 6, 11, 9, 12, 12]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C5"
    if result["warnings"] or result.get("notes"):
        notes = wb.create_sheet("Checks")
        for w in result["warnings"] + result.get("notes", []):
            notes.append([w])
    about = wb.create_sheet("About")
    about.append(["Durations are provisional planning assumptions until confirmed by the site team."])
    about.append(["Made with the construction-critical-path skill. Full workflow: https://constructions.bleno.io · pilots: https://bleno.io"])
    wb.save(path)


def _workdays(spec):
    days = set()
    for part in spec.split(","):
        if "-" in part:
            a, b = part.split("-")
            days.update(range(int(a), int(b) + 1))
        elif part.strip():
            days.add(int(part))
    if not days or not days <= set(range(1, 8)):
        raise argparse.ArgumentTypeError("workdays are ISO weekday numbers 1 (Mon) … 7 (Sun), e.g. 1-6 or 1-5")
    return days


def main(argv=None):
    ap = argparse.ArgumentParser(description="Critical-path (CPM) calculator for construction schedules.")
    ap.add_argument("schedule", help="JSON or CSV schedule")
    ap.add_argument("--json", action="store_true", help="print the full result as JSON")
    ap.add_argument("--csv", metavar="FILE", help="write the result as CSV")
    ap.add_argument("--xlsx", metavar="FILE", help="write an Excel workbook with a bar chart (needs openpyxl)")
    ap.add_argument("--start", metavar="YYYY-MM-DD", help="commencement date, to show calendar dates")
    ap.add_argument("--workdays", type=_workdays, default=_workdays("1-6"), help="working weekdays, default 1-6 (Mon–Sat)")
    ap.add_argument("--holiday", action="append", default=[], metavar="YYYY-MM-DD", help="a non-working date (repeatable)")
    ap.add_argument("--deadline", metavar="YYYY-MM-DD", help="contractual completion date to compare with (needs --start)")
    args = ap.parse_args(argv)

    project, acts = load(args.schedule)
    try:
        result = calculate(acts)
    except ValueError as e:
        print(f"Cannot calculate: {e}", file=sys.stderr)
        return 2

    dates, day_labels = {}, None
    result["notes"] = []
    if args.deadline and not args.start:
        print("--deadline needs --start", file=sys.stderr)
        return 2
    if args.start:
        start = dt.date.fromisoformat(args.start)
        holidays = set(args.holiday)
        if not _working(start, args.workdays, holidays):
            first = to_date(0, start, args.workdays, holidays, "start")
            result["notes"].append(f"{start.isoformat()} is not a working day; work starts on {first.isoformat()}.")
        for r in result["rows"]:
            s = to_date(r["es"], start, args.workdays, holidays, "start")
            f = to_date(r["ef"], start, args.workdays, holidays, "finish")
            r["start"], r["finish"] = s.isoformat(), f.isoformat()
            dates[r["id"]] = (r["start"], r["finish"])
        result["finish_date"] = to_date(result["finish_days"], start, args.workdays, holidays, "finish").isoformat()
        day_labels, d = [], to_date(0, start, args.workdays, holidays, "start")
        while len(day_labels) < int(result["finish_days"]):
            if _working(d, args.workdays, holidays):
                day_labels.append(d.strftime("%d %b"))
            d += dt.timedelta(days=1)
        if args.deadline:
            deadline = dt.date.fromisoformat(args.deadline)
            finish_date = dt.date.fromisoformat(result["finish_date"])
            gap = (deadline - finish_date).days
            result["deadline"] = {"date": deadline.isoformat(), "calendar_days_to_spare": gap}
            result["notes"].append(
                f"Completion {finish_date.isoformat()} is {abs(gap)} calendar days "
                f"{'before' if gap >= 0 else 'AFTER'} the deadline {deadline.isoformat()}.")

    if args.csv:
        write_csv(args.csv, result, dates)
    if args.xlsx:
        write_xlsx(args.xlsx, project, result, dates, day_labels)
    if args.json:
        print(json.dumps({"project": project, **result}, indent=2, ensure_ascii=False))
    else:
        print(table(project, result, dates))
        if result.get("finish_date"):
            print(f"Completion: {result['finish_date']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
