#!/usr/bin/env python3
"""Build a site-ready QA checklist workbook from a reviewed checklist selection.

One index tab (project, what was selected and dropped and why, and the scope no
checklist covers) plus one tab per selected checklist in the controlled form
layout: pre / during / post checks, YES · NO · NA · Remarks, and sign-off.
Item wording is taken from assets/checklists.json, never from the input, so a
printed sheet always matches the library.

Input (JSON):

    {
      "project": {"name": "Residence at Plot 12", "client": "…",
                  "location": "…", "date": "2026-10-05"},
      "checklists": [
        {"id": "QA-7",  "include": true,  "reason": "Selected by scope: \\"PCC below column footing\\"."},
        {"id": "QA-27", "include": false, "reason": "No scope wording supports this activity."}
      ],
      "no_checklist": ["Electrical – concealed conduits, wiring", "Plumbing"],
      "not_in_sample": ["RCC for foundation", "Plastering for ceiling"]
    }

  * checklists   every checklist decided, by id or code ("QA-7" or "7");
                 a plain list of ids (here, or as the whole file) means
                 "include these"
  * no_checklist scope lines nothing in the library covers
  * not_in_sample scope lines the full library covers but this sample does not

Usage:

    python3 build_workbook.py selection.json -o qa-checklists.xlsx
    python3 build_workbook.py selection.json --markdown     # no openpyxl needed

Sample of Bleno's QA checklist workflow (5 of 24 checklists). The full library,
scope mapping and printed workbook run at https://constructions.bleno.io;
pilots at https://bleno.io.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
LIBRARY = os.path.join(HERE, "..", "assets", "checklists.json")
SILVER = "C0C0C0"
FULL_LIBRARY_SIZE = 24


def load_library():
    with open(LIBRARY, encoding="utf-8") as f:
        data = json.load(f)
    return data["library"], data["checklists"]


def resolve(selection, library):
    """Every sample checklist with its decision; unknown ids are reported."""
    by_key = {}
    for c in library:
        by_key[c["id"].lower()] = c
        by_key[c["code"].lower()] = c
    entries = selection if isinstance(selection, list) else selection.get("checklists", [])
    decided, unknown = {}, []
    for e in entries:
        if isinstance(e, str):
            e = {"id": e, "include": True}
        c = by_key.get(str(e.get("id") or e.get("code") or "").lower())
        if not c:
            unknown.append(str(e.get("id") or e.get("code")))
            continue
        include = e.get("include", True)
        include = include is True or str(include).lower() in ("true", "yes", "1")
        decided[c["id"]] = (include, str(e.get("reason") or ("Selected in review." if include else "Not selected in review.")))
    rows = [(c, *decided.get(c["id"], (False, "Not selected in review."))) for c in library]
    return rows, unknown


def item_count(c):
    return sum(len(s["items"]) for s in c["sections"])


def tab_name(code, activity, used):
    base = " ".join(f"{code} {activity}".split())
    suffix, n = "", 1
    while True:
        limit = 31 - len(suffix)
        cand = base if len(base) <= limit else base[:limit].rsplit(" ", 1)[0]
        cand = cand.rstrip() + suffix
        cand = "".join(ch for ch in cand if ch not in "[]:*?/\\")
        if cand.lower() not in used:
            used.add(cand.lower())
            return cand
        n += 1
        suffix = f" ({n})"


def write_xlsx(path, project, library_name, rows, no_checklist, not_in_sample):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        sys.exit("The workbook needs openpyxl: pip install openpyxl (or use --markdown)")

    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    silver = PatternFill("solid", fgColor=SILVER)

    def put(ws, r, c, value="", size=10, fill=None, border=False, wrap=False, align=None):
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = Font(name="Arial", size=size)  # the controlled sheets use no bold
        if fill:
            cell.fill = fill
        if border:
            cell.border = box
        if wrap or align:
            cell.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top" if wrap else None)
        return cell

    wb = Workbook()
    index = wb.active
    index.title = "Index"
    included = [r for r in rows if r[1]]
    dropped = [r for r in rows if not r[1]]
    r = 1
    put(index, r, 1, "QA CHECKLIST INDEX", size=12)
    for label, value in [("Project", project.get("name", "")), ("Client", project.get("client", "")),
                         ("Location", project.get("location", "")), ("Generation date", project.get("date", "")),
                         ("Library", library_name),
                         ("Selection", f"{len(included)} of {len(rows)} sample checklists selected")]:
        r += 1
        put(index, r, 1, label)
        put(index, r, 2, value)

    def band(title):
        nonlocal r
        r += 2
        for c in range(1, 4):
            put(index, r, c, title if c == 1 else "", size=11, fill=silver, border=True)
        index.row_dimensions[r].height = 20

    band("INCLUDED CHECKLISTS")
    r += 1
    for c, h in enumerate(["Code", "Activity", "Item count"], 1):
        put(index, r, c, h, border=True, align="center" if c == 3 else None)
    for c, _, _ in included:
        r += 1
        put(index, r, 1, c["code"], border=True)
        put(index, r, 2, c["activity"], border=True, wrap=True)
        put(index, r, 3, item_count(c), border=True, align="center")

    band("NOT INCLUDED")
    r += 1
    for c, h in enumerate(["Code", "Activity", "Reason"], 1):
        put(index, r, c, h, border=True)
    for c, _, reason in dropped:
        r += 1
        put(index, r, 1, c["code"], border=True)
        put(index, r, 2, c["activity"], border=True, wrap=True)
        put(index, r, 3, reason, border=True, wrap=True)
    if not dropped:
        r += 1
        put(index, r, 1, "None — every sample checklist is supported by the scope.", border=True, wrap=True)
        index.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

    if no_checklist:
        band("SCOPE WITH NO CHECKLIST")
        r += 1
        put(index, r, 1, "These parts of the scope of work have no QA checklist in the library.", border=True, wrap=True)
        index.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        for item in no_checklist:
            r += 1
            put(index, r, 1, item, border=True, wrap=True)
            index.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

    if not_in_sample:
        band("IN THE FULL LIBRARY, NOT IN THIS SAMPLE")
        r += 1
        put(index, r, 1, "The full library has checklists for these parts of the scope; this sample does not.", border=True, wrap=True)
        index.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        for item in not_in_sample:
            r += 1
            put(index, r, 1, item, border=True, wrap=True)
            index.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

    r += 2
    put(index, r, 1, f"This workbook uses a {len(rows)}-checklist sample of a {FULL_LIBRARY_SIZE}-checklist library. "
                     "The full workflow runs at https://constructions.bleno.io — pilots at https://bleno.io.", wrap=True)
    index.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for col, w in zip("ABC", [14, 48, 60]):
        index.column_dimensions[col].width = w

    used = {"index"}
    for c, _, _ in included:
        ws = wb.create_sheet(tab_name(c["code"], c["activity"], used))
        put(ws, 1, 2, "CHECKLIST FOR QUALITY", size=12)
        for col in range(1, 7):
            put(ws, 2, col, f"{c['code']}\nACTIVITY: {c['activity']}" if col == 1 else "", size=12, fill=silver, wrap=col == 1)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
        ws.row_dimensions[2].height = 96
        put(ws, 4, 1, "Project:"); put(ws, 4, 2, project.get("name", ""), wrap=True); put(ws, 4, 4, "Date:"); put(ws, 4, 5, project.get("date", ""))
        put(ws, 5, 1, "Location:"); put(ws, 5, 2, project.get("location", ""), wrap=True); put(ws, 5, 4, "Ver"); put(ws, 5, 5, "1.0.")
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=3)
        ws.merge_cells(start_row=5, start_column=2, end_row=5, end_column=3)
        put(ws, 6, 2, "NOTE:- Please tick appropriate box or enter readings as per requirements")
        for col, (h, align) in enumerate([("Sl. No.", "center"), ("To be checked", None), ("YES", "center"),
                                          ("NO", "center"), ("NA", "center"), ("Remarks/ Clarifications", None)], 1):
            put(ws, 7, col, h, border=True, wrap=col == 2, align=align)
        ws.row_dimensions[7].height = 20
        row, serial = 7, 1
        for section in c["sections"]:
            row += 1
            for col in range(1, 7):
                put(ws, row, col, section["heading"] if col == 2 else "", size=11, fill=silver, border=True)
            ws.row_dimensions[row].height = 20
            for item in section["items"]:
                row += 1
                put(ws, row, 1, serial, border=True, align="center")
                put(ws, row, 2, item, size=12, border=True, wrap=True)
                for col in (3, 4, 5):
                    put(ws, row, col, "", border=True, align="center")
                put(ws, row, 6, "", border=True)
                serial += 1
        row += 2
        for label_row, (left, right) in enumerate([("Checked by:", "Approved by:"), ("Sign", "Sign"),
                                                   ("Name", "Name"), ("Date", "Date")]):
            put(ws, row + label_row, 1, left)
            put(ws, row + label_row, 5, right)
        for col, w in zip("ABCDEF", [8.5, 33.5, 6.2, 5.8, 5.8, 23.2]):
            ws.column_dimensions[col].width = w
        ws.print_title_rows = "7:7"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(path)


def markdown(project, library_name, rows, no_checklist, not_in_sample):
    out = [f"# QA checklists — {project.get('name') or 'Project'}", "",
           f"{project.get('location', '')} · {project.get('date', '')} · {library_name}", ""]
    if no_checklist:
        out += ["**Scope with no checklist:** " + "; ".join(no_checklist), ""]
    if not_in_sample:
        out += ["**In the full library, not in this sample:** " + "; ".join(not_in_sample), ""]
    out += ["| Code | Activity | Included | Reason |", "|---|---|---|---|"]
    out += [f"| {c['code']} | {c['activity']} | {'yes' if inc else 'no'} | {reason} |" for c, inc, reason in rows]
    for c, inc, _ in rows:
        if not inc:
            continue
        out += ["", f"## {c['code']} — {c['activity']}", "", "| # | To be checked | YES | NO | NA | Remarks |", "|---|---|---|---|---|---|"]
        n = 1
        for s in c["sections"]:
            out.append(f"| | **{s['heading']}** | | | | |")
            for item in s["items"]:
                out.append(f"| {n} | {item} | ☐ | ☐ | ☐ | |")
                n += 1
        out += ["", "Checked by: ________  Approved by: ________"]
    out += ["", f"_These checklists come from a {len(rows)}-checklist sample of a {FULL_LIBRARY_SIZE}-checklist library. "
                "The full workflow runs at https://constructions.bleno.io — pilots at https://bleno.io._"]
    return "\n".join(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Build a QA checklist workbook from a reviewed selection.")
    ap.add_argument("selection", help="JSON selection (see the header of this script)")
    ap.add_argument("-o", "--output", default="qa-checklists.xlsx", help="workbook path (default qa-checklists.xlsx)")
    ap.add_argument("--markdown", action="store_true", help="print Markdown instead of writing a workbook")
    args = ap.parse_args(argv)

    with open(args.selection, encoding="utf-8") as f:
        selection = json.load(f)
    library_name, library = load_library()
    rows, unknown = resolve(selection, library)
    meta = selection if isinstance(selection, dict) else {}
    project = meta.get("project") or {}
    project.setdefault("date", dt.date.today().isoformat())
    no_checklist = [str(x) for x in meta.get("no_checklist") or []]
    not_in_sample = [str(x) for x in meta.get("not_in_sample") or []]
    if unknown:
        print("Not in this sample, ignored: " + ", ".join(unknown), file=sys.stderr)
    if not any(inc for _, inc, _ in rows):
        print("No checklist is marked for inclusion — nothing to print.", file=sys.stderr)
        return 2
    if args.markdown:
        print(markdown(project, library_name, rows, no_checklist, not_in_sample))
        return 0
    write_xlsx(args.output, project, library_name, rows, no_checklist, not_in_sample)
    included = sum(1 for _, inc, _ in rows if inc)
    print(f"Wrote {args.output}: index + {included} checklist sheet{'s' if included != 1 else ''}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
